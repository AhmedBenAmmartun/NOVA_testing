from __future__ import annotations

import re
from pathlib import Path


_HEADING = re.compile(r"^#{1,6}\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")


def parse_outline(markdown: str) -> list[tuple[str, list[str]]]:
    slides: list[tuple[str, list[str]]] = []
    title: str | None = None
    bullets: list[str] = []

    def flush() -> None:
        nonlocal title, bullets
        if title:
            slides.append((title.strip(), [item.strip() for item in bullets if item.strip()]))
        title = None
        bullets = []

    for raw in (markdown or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = _HEADING.match(line)
        if heading:
            flush()
            title = heading.group(1).strip()
            continue
        bullet = _BULLET.match(line) or _NUMBERED.match(line)
        if bullet:
            if title is None:
                title = "Key Points"
            bullets.append(bullet.group(1).strip())
            continue
        if title is None:
            title = line[:80]
        else:
            bullets.append(line)

    flush()
    return slides


def build_presentation(
    outline_markdown: str,
    destination: Path,
    *,
    course: str,
    title: str,
    subtitle: str = "NOVA Class Intelligence recap",
) -> Path:
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except Exception as error:
        raise RuntimeError(
            "python-pptx is required to build the class recap presentation."
        ) from error

    destination = destination.expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)

    deck = Presentation()
    title_slide = deck.slides.add_slide(deck.slide_layouts[0])
    title_slide.shapes.title.text = title
    title_slide.placeholders[1].text = f"{course}\n{subtitle}"

    slides = parse_outline(outline_markdown)
    if not slides:
        slides = [("Session Recap", ["See the accompanying class notes for details."])]

    for slide_title, bullets in slides[:18]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = slide_title[:120]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, bullet in enumerate(bullets[:7] or ["See class notes for details."]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet[:500]
            paragraph.level = 0
            paragraph.font.size = Pt(22)

    deck.save(str(destination))
    return destination
